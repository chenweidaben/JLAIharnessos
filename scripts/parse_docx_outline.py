# -*- coding: utf-8 -*-
"""解析全诊产品说明.docx - 获取大纲结构和基本统计"""
import docx
import json
import os

DOC_PATH = r"C:\claudecode5.1.118\全诊产品说明.docx"
OUT_DIR = r"C:\Users\lenovo\Doubao\chats\2026-09-14\new-chat-1\docs\requirements"
os.makedirs(OUT_DIR, exist_ok=True)

doc = docx.Document(DOC_PATH)

# 统计
para_count = len(doc.paragraphs)
table_count = len(doc.tables)
section_count = len(doc.sections)

print(f"=== 文档基本统计 ===")
print(f"段落数: {para_count}")
print(f"表格数: {table_count}")
print(f"节数: {section_count}")

# 提取大纲（标题样式）
print(f"\n=== 文档大纲（标题层级）===")
outline = []
for i, para in enumerate(doc.paragraphs):
    style_name = para.style.name if para.style else ""
    text = para.text.strip()
    if not text:
        continue
    # 检测标题样式
    if style_name.startswith("Heading") or style_name.startswith("标题") or "标题" in style_name:
        level = ""
        if style_name.startswith("Heading"):
            try:
                level = style_name.replace("Heading ", "").strip()
            except:
                level = "?"
        elif style_name.startswith("标题"):
            level = style_name.replace("标题", "").strip() or "1"
        indent = "  " * (int(level) - 1) if level.isdigit() else ""
        print(f"{indent}[{style_name}|L{level}] {text[:80]}")
        outline.append({"index": i, "style": style_name, "level": level, "text": text})

# 也检测非标题样式但看起来像标题的段落（编号开头）
print(f"\n=== 疑似标题（编号开头的段落，前100条）===")
import re
suspected = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    style_name = para.style.name if para.style else ""
    # 跳过已识别的标题
    if style_name.startswith("Heading") or style_name.startswith("标题"):
        continue
    # 匹配编号模式：1. / 1.1 / 一、 / （一） / 第X章
    if re.match(r'^(\d+\.\d*|[一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|第[一二三四五六七八九十\d]+[章节篇])', text):
        if len(suspected) < 100:
            print(f"  [{i}|{style_name}] {text[:80]}")
        suspected.append({"index": i, "style": style_name, "text": text})

print(f"\n疑似标题总数: {len(suspected)}")

# 保存大纲到JSON
with open(os.path.join(OUT_DIR, "_outline.json"), "w", encoding="utf-8") as f:
    json.dump({"outline": outline, "suspected_headings": suspected, 
               "stats": {"paragraphs": para_count, "tables": table_count}}, 
              f, ensure_ascii=False, indent=2)

print(f"\n大纲已保存到: {os.path.join(OUT_DIR, '_outline.json')}")
